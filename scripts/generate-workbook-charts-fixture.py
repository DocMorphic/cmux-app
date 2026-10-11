#!/usr/bin/env python3
"""Self-authored worksheet charts with intentionally absent chart caches."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from datetime import datetime
import xml.etree.ElementTree as ET
import re
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, DoughnutChart, RadarChart, ScatterChart, BubbleChart, AreaChart, PieChart, Reference, Series

root = Path(__file__).resolve().parents[1]
book = Workbook(); sheet = book.active; sheet.title = 'Charts'
data = book.create_sheet("Data's résumé")
for row in [('Quarter', 'Revenue', 'Plan', 'Size'), ('Q1 中', 100, 80, 10), ('Q2', 175, 160, 15), ('Q3', 125, 110, 12)]:
    data.append(row)
sheet['A1'] = 'CMUX WORKBOOK CHARTS'
for name in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']: sheet.column_dimensions[name].width = 18
for row in range(1, 55): sheet.row_dimensions[row].height = 24

def category(cls, title, position, worksheet=sheet):
    chart = cls(); chart.title = title; chart.width = 12; chart.height = 7
    chart.add_data(Reference(data, min_col=2, max_col=3 if cls in [LineChart, AreaChart, RadarChart] else 2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=4))
    worksheet.add_chart(chart, position)
    return chart

category(BarChart, 'Column chart', 'A4')
category(LineChart, 'Line chart', 'F4')
category(DoughnutChart, 'Doughnut chart', 'A19')
category(RadarChart, 'Radar chart', 'F19')
category(AreaChart, 'Area chart', 'A34')
category(PieChart, 'Pie chart', 'F34')
scatter = ScatterChart(); scatter.title = 'Far scatter chart'; scatter.width = 12; scatter.height = 7
scatter.series.append(Series(Reference(data, min_col=3, min_row=2, max_row=4), Reference(data, min_col=2, min_row=2, max_row=4), title='Plan against revenue'))
sheet.add_chart(scatter, 'A131')
bubble = BubbleChart(); bubble.title = 'Bubble chart'; bubble.width = 12; bubble.height = 7
bubble.series.append(Series(Reference(data, min_col=3, min_row=2, max_row=4), Reference(data, min_col=2, min_row=2, max_row=4), Reference(data, min_col=4, min_row=2, max_row=4), title='Sized plan'))
sheet.add_chart(bubble, 'F131')
only = book.create_sheet('Chart only')
category(BarChart, 'Chart-only worksheet', 'B12', only)
book.properties.created = book.properties.modified = datetime(2026, 10, 10)
raw = BytesIO(); book.save(raw)
with ZipFile(raw) as archive: parts = {name: archive.read(name) for name in archive.namelist()}
# openpyxl overwrites modified time during save; normalize the actual package.
parts['docProps/core.xml'] = re.sub(rb'(<dcterms:modified[^>]*>)[^<]+',
    rb'\g<1>2026-10-10T00:00:00Z', parts['docProps/core.xml'])
xdr = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
labels = [['Column chart', 'Line chart', 'Doughnut chart', 'Radar chart', 'Area chart', 'Pie chart', 'Far scatter chart', 'Bubble chart'], ['Chart-only worksheet']]
for index, names in enumerate(labels, 1):
    path = f'xl/drawings/drawing{index}.xml'; node = ET.fromstring(parts[path])
    for properties, name in zip(node.findall(f'.//{{{xdr}}}cNvPr'), names): properties.set('descr', name)
    parts[path] = ET.tostring(node, encoding='utf-8', xml_declaration=True)
output = root / 'app/src/androidTest/assets/workbook/charts.xlsx'
with ZipFile(output, 'w') as archive:
    for name, content in parts.items():
        info = ZipInfo(name, (2026, 10, 10, 0, 0, 0)); info.compress_type = ZIP_DEFLATED; info.external_attr = 0o600 << 16
        archive.writestr(info, content)
print(output)
