#!/usr/bin/env python3
"""Generate a small, self-authored PowerPoint preview fixture (python-pptx/Pillow)."""
from pathlib import Path
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image

root = Path(__file__).resolve().parents[1]
result = root / 'app/src/androidTest/assets/presentation/rich.pptx'
result.parent.mkdir(parents=True, exist_ok=True)
p = Presentation(); p.slide_width = Inches(10); p.slide_height = Inches(7.5)
slides = [p.slides.add_slide(p.slide_layouts[6]) for _ in range(3)]
def text(slide, value, x, y, w=8, h=1, size=28):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = shape.text_frame.paragraphs[0].add_run(); run.text = value
    run.font.size = Pt(size); run.font.color.rgb = RGBColor(20, 30, 45)
    return shape, run
text(slides[0], 'CMUX POWERPOINT', 0.5, 0.3)
text(slides[0], 'Offline slides · Résumé 日本語', 0.5, 1.1, size=22)
png = BytesIO(); Image.new('RGB', (160, 120), (235, 30, 35)).save(png, format='PNG'); png.seek(0)
slides[0].shapes.add_picture(png, Inches(0.8), Inches(2.3), width=Inches(3.2), height=Inches(2.4))
link, _ = text(slides[0], 'Jump to final slide', 0.5, 5.2); link.click_action.target_slide = slides[2]
_, run = text(slides[0], 'Blocked active link', 0.5, 6.2, size=18)
run.hyperlink.address = 'javascript:window.__cmuxUnsafeLink=true'
text(slides[1], 'Revenue and plan', 0.5, 0.3)
table = slides[1].shapes.add_table(3, 2, Inches(.5), Inches(1.2), Inches(4), Inches(1.8)).table
for r, values in enumerate([['Quarter', 'Revenue'], ['Q1', '$100'], ['Q2', '$175']]):
    for c, value in enumerate(values): table.cell(r,c).text = value
chart = CategoryChartData(); chart.categories = ['Q1', 'Q2']; chart.add_series('Revenue', [100, 175])
slides[1].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.5), Inches(3.2), Inches(8.7), Inches(3.8), chart)
text(slides[2], 'Final slide 你好', .5, .3)
text(slides[2], 'Selected slide survives restoration.', .5, 1.4, size=22)
link, _ = text(slides[2], 'Back to cover', .5, 3); link.click_action.target_slide = slides[0]
p.save(result)
print(result)
